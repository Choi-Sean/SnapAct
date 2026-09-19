#!/usr/bin/env python3
"""docs/SnapAct_클래스별액션_검토.xlsx -> Resources/actions.json

The spreadsheet is the source of truth for classes, actions, priors and
signals. Swift only parses the JSON this writes; nothing in the package
hardcodes a class name, a verb, a score or a weight. Editing the sheet and
re-running this script is the only way the catalog changes.

    python build_catalog.py            # generate
    python build_catalog.py --check    # fail if the sheet is newer than the JSON
    python build_catalog.py --dry-run  # report only, write nothing

Verb extraction is deliberately conservative. The 주/부 액션 columns are
Korean display text that usually carries the verb in parentheses — "연락처 추가
(create_contact)" — but often does not. An unparsed cell is reported as
unmapped, never guessed: a wrong verb here becomes a button that does the
wrong thing to someone's calendar.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
import sys
from datetime import datetime, timezone
from pathlib import Path

import openpyxl

REPO = Path(__file__).resolve().parents[3]
XLSX = REPO / "docs/SnapAct_클래스별액션_검토.xlsx"
RESOURCES = REPO / "packages/SnapActKit/Sources/SnapActKit/Resources"
OUT_JSON = RESOURCES / "actions.json"
UNMAPPED_CSV = Path(__file__).resolve().parent / "unmapped.csv"

SCHEMA_VERSION = 1

# Base scores before any personalisation. Ranking multiplies these; it never
# reorders by them alone.
SCORE_PRIMARY = 0.70
SCORE_SECONDARY_FIRST = 0.45
SCORE_SECONDARY_REST = 0.35
SCORE_UNIVERSAL = 0.20

# From the 읽어보기 sheet: "모든 클래스에 유니버설 액션(노트 저장·리마인더·복사·
# 검색·번역)과 [기타] 버튼이 항상 함께 깔립니다." It is prose in a readme sheet
# rather than a column, so it is restated here rather than parsed — but it is
# still emitted into the JSON, so Swift does not hardcode it either.
UNIVERSAL_VERBS = ["save_note", "create_reminder", "copy_text", "search", "translate"]

# Display text -> verb, inside parentheses. Trailing qualifiers are ignored:
# "(create_reminder, 기한 3일 전)" and "(create_event 다중)" both yield the verb.
VERB_IN_PARENS = re.compile(r"\(([a-z_][a-z0-9_]*)\b[^)]*\)")
SECONDARY_SPLIT = re.compile(r"\s*·\s*")

CONFIRMATION = {"자동": "auto", "되돌림 가능": "reversible", "명시적 확인": "explicit"}


# ---------------------------------------------------------------------------
# Review adjustments — from the 수정의견 column (9 entries, 6 decisions).
#
# These live here rather than in the sheet because the 수정의견 cells are free
# prose that cannot be parsed into structure. Each carries the opinion it
# implements so the decision stays auditable against the source.
# ---------------------------------------------------------------------------

# (1) book_cover / receipt / nutrition_label — "앨범 생성 및 캡션 추가"
#
# Album creation is real: PHAssetCollectionChangeRequest can create an album
# and add assets, so `add_to_album` is a genuine verb.
#
# CAPTIONS ARE NOT POSSIBLE. PhotoKit exposes no API for writing the Photos
# caption field; a third-party app cannot set it. There is deliberately NO
# caption verb in this catalog — inventing one would produce a button that
# cannot work. The substitute is save_note into our own store, linked to the
# asset by PHAsset.localIdentifier, which survives and is searchable.
ALBUM_CLASSES = ["book_cover", "receipt", "nutrition_label"]
ADD_TO_ALBUM = {
    "verb": "add_to_album",
    "api": "PhotoKit (PHAssetCollectionChangeRequest)",
    "confirmation": "reversible",
    "undoable": True,
    "note": "앨범 생성·추가는 가능. 사진 캡션 쓰기는 PhotoKit에 없어 save_note로 대체한다.",
}

# (2) device_display / workout_display / nutrition_label — "health related class"
# Grouped, NOT merged: their tiers differ (device_display=0, nutrition_label=2),
# so merging would drag a Tier 2 nutrition label into Tier 0 handling.
HEALTH_GROUP = ["device_display", "workout_display", "nutrition_label"]

# (3) payment_screenshot — "receipt랑 똑같이 처리"
ALIAS_OF = {"payment_screenshot": "receipt"}
# ...but the masking rule stays: a payment screenshot can show an account number.
MASKING_RULES = {"payment_screenshot": ["account_number"]}

# (4) error_screen — "llm 에서 처리 필 클래스 따로 분류"
LLM_RESOLUTION = [
    "error_screen", "handwritten_note", "chat_screenshot",
    "event_flyer", "parking_restriction",
]

# (5) appointment_slip — "스크린샷의 경우가 더 많을 듯"
# Class kept. When the screenshot signal fires, both candidates are offered
# rather than one replacing the other — ranking reorders, it never removes.
CO_PRESENT_WHEN_SCREENSHOT = {"appointment_slip": ["booking_screenshot"]}

# (6) chat_screenshot — "캘린더 또는 노트 상 어떠한 정보를 기억"
CHAT_PRIMARY = "create_event"
CHAT_SECONDARY = ["save_note", "create_reminder"]
# Third-party data: only the extracted appointment survives, never the
# conversation text. Enforced downstream by the logger, flagged here so the
# rule travels with the class.
NO_RAW_RETENTION = ["chat_screenshot"]


def _cell(row, i):
    return str(row[i]).strip() if i < len(row) and row[i] is not None else ""


def load_verbs(wb) -> dict:
    verbs = {}
    for row in list(wb["액션어휘"].iter_rows(values_only=True))[1:]:
        name = _cell(row, 0)
        if not re.fullmatch(r"[a-z_][a-z0-9_]*", name):
            continue  # the sheet's trailing prose rows ("에이전트 규칙", "· ...")
        verbs[name] = {
            "api": _cell(row, 1),
            "confirmation": CONFIRMATION.get(_cell(row, 2), _cell(row, 2)),
            "undoable": _cell(row, 3) == "가능",
            "note": _cell(row, 4),
        }
    return verbs


def extract(text: str, whitelist: set[str], *, split: bool) -> tuple[list[dict], list[str]]:
    """Returns (mapped, unmapped_display_text) for one cell.

    `split` only for 부 액션, which is a "·"-separated list. A 주 액션 cell is
    ONE action whose display text may itself contain "·" — splitting it
    shredded transit_sign's "노선·도착 조회 (search ...)" into a bogus unmapped
    "노선" plus a mapped remainder.
    """
    mapped, unmapped = [], []
    parts = SECONDARY_SPLIT.split(text) if split else ([text] if text.strip() else [])
    for part in (p for p in parts if p.strip()):
        found = [v for v in VERB_IN_PARENS.findall(part) if v in whitelist]
        if found:
            display = VERB_IN_PARENS.sub("", part).strip(" ·")
            for v in found:
                mapped.append({"verb": v, "display": display or part.strip()})
        else:
            unmapped.append(part.strip())
    return mapped, unmapped


def load_classes(wb, whitelist: set[str]) -> tuple[dict, list[dict]]:
    classes, unmapped_report = {}, []
    for row in list(wb["클래스별액션"].iter_rows(values_only=True))[1:]:
        name = _cell(row, 1)
        if not name:
            continue

        primary, un_primary = extract(_cell(row, 8), whitelist, split=False)
        secondary, un_secondary = extract(_cell(row, 9), whitelist, split=True)

        for slot, items in (("primary", un_primary), ("secondary", un_secondary)):
            for display in items:
                unmapped_report.append({"class": name, "slot": slot, "display": display})

        for i, a in enumerate(primary):
            a["baseScore"] = SCORE_PRIMARY
        for i, a in enumerate(secondary):
            a["baseScore"] = SCORE_SECONDARY_FIRST if i == 0 else SCORE_SECONDARY_REST

        classes[name] = {
            "id": int(_cell(row, 0)) if _cell(row, 0).isdigit() else None,
            "motive": _cell(row, 2),
            "situation": _cell(row, 3),
            "priority": _cell(row, 4),
            "tier": int(_cell(row, 5)) if _cell(row, 5).isdigit() else None,
            "terminatingLayer": _cell(row, 6),
            "extractionFields": [f.strip() for f in _cell(row, 7).split(",") if f.strip()],
            "primary": primary,
            "secondary": secondary,
            "unmapped": [
                {"slot": s, "display": d}
                for s, items in (("primary", un_primary), ("secondary", un_secondary))
                for d in items
            ],
            "fallback": _cell(row, 10),
            "promotionBasis": _cell(row, 11),
            "pitfalls": _cell(row, 12),
            # Filled by apply_review_adjustments.
            "group": None,
            "aliasOf": None,
            "resolution": "rules",
            "maskingRules": [],
            "coPresentWhenScreenshot": [],
            "retainsRawText": True,
        }
    return classes, unmapped_report


def load_blocking(wb) -> dict:
    out = {}
    for row in list(wb["차단클래스"].iter_rows(values_only=True))[1:]:
        name = _cell(row, 1)
        if not re.fullmatch(r"[a-z_]+", name):
            continue  # trailing "공통 원칙" prose rows
        out[name] = {
            "id": int(_cell(row, 0)) if _cell(row, 0).isdigit() else None,
            "visualDiscriminators": _cell(row, 2),
            "userFacing": _cell(row, 3),
            "provides": _cell(row, 4),
            "confusionRisk": _cell(row, 5),
            "tier": 0,
        }
    return out


def load_signals(wb) -> list[dict]:
    out = []
    for row in list(wb["랭킹신호"].iter_rows(values_only=True))[1:]:
        name, source = _cell(row, 0), _cell(row, 1)
        if not name or not source:
            continue  # trailing "핵심 원칙" prose rows have no source
        out.append({
            "signal": name,
            "source": source,
            "inference": _cell(row, 2),
            "affects": _cell(row, 3),
        })
    return out


def _clear_unmapped(cls: dict, slot: str) -> None:
    """An adjustment that replaces a slot also settles its unmapped entries.

    Leaving them behind would report "일정 또는 알림 생성 is unmapped" for a
    class whose primary action the review already decided, sending someone to
    map a cell that no longer drives anything.
    """
    cls["unmapped"] = [u for u in cls["unmapped"] if u["slot"] != slot]


def apply_review_adjustments(classes: dict, verbs: dict, warn) -> None:
    verbs[ADD_TO_ALBUM["verb"]] = {k: v for k, v in ADD_TO_ALBUM.items() if k != "verb"}

    for name in ALBUM_CLASSES:
        if name not in classes:
            warn(f"(1) 앨범: 클래스 없음 {name}")
            continue
        sec = classes[name]["secondary"]
        if not any(a["verb"] == "add_to_album" for a in sec):
            sec.append({"verb": "add_to_album", "display": "앨범에 추가",
                        "baseScore": SCORE_SECONDARY_REST})

    for name in HEALTH_GROUP:
        if name in classes:
            classes[name]["group"] = "health"
        else:
            warn(f"(2) health 그룹: 클래스 없음 {name}")

    for alias, target in ALIAS_OF.items():
        if alias not in classes or target not in classes:
            warn(f"(3) alias: {alias} -> {target} 중 하나가 없음")
            continue
        classes[alias]["aliasOf"] = target
        # Materialised here so Swift stays a pure parser and the debug screen
        # can show the inherited list without resolving anything itself.
        classes[alias]["primary"] = [dict(a) for a in classes[target]["primary"]]
        classes[alias]["secondary"] = [dict(a) for a in classes[target]["secondary"]]
        classes[alias]["unmapped"] = [dict(u) for u in classes[target]["unmapped"]]

    for name, rules in MASKING_RULES.items():
        if name in classes:
            classes[name]["maskingRules"] = list(rules)

    for name in LLM_RESOLUTION:
        if name in classes:
            classes[name]["resolution"] = "llm"
        else:
            warn(f"(4) llm resolution: 클래스 없음 {name}")

    for name, others in CO_PRESENT_WHEN_SCREENSHOT.items():
        if name in classes:
            classes[name]["coPresentWhenScreenshot"] = list(others)

    if "chat_screenshot" in classes:
        c = classes["chat_screenshot"]
        c["primary"] = [{"verb": CHAT_PRIMARY, "display": "캘린더 일정",
                         "baseScore": SCORE_PRIMARY}]
        c["secondary"] = [
            {"verb": v, "display": d,
             "baseScore": SCORE_SECONDARY_FIRST if i == 0 else SCORE_SECONDARY_REST}
            for i, (v, d) in enumerate(zip(CHAT_SECONDARY, ["노트 저장", "알림 생성"]))
        ]
        _clear_unmapped(c, "primary")
        _clear_unmapped(c, "secondary")
    else:
        warn("(6) chat_screenshot 클래스 없음")

    for name in NO_RAW_RETENTION:
        if name in classes:
            classes[name]["retainsRawText"] = False


def validate(classes: dict, verbs: dict, blocking: dict) -> list[str]:
    errors = []
    whitelist = set(verbs)
    for name, c in classes.items():
        for slot in ("primary", "secondary"):
            for a in c[slot]:
                if a["verb"] not in whitelist:
                    errors.append(f"{name}.{slot}: 어휘에 없는 동사 {a['verb']!r}")
    for v in UNIVERSAL_VERBS:
        if v not in whitelist:
            errors.append(f"유니버설 동사가 어휘에 없음: {v!r}")
    overlap = set(classes) & set(blocking)
    if overlap:
        errors.append(f"서비스 클래스와 차단 클래스가 겹침: {sorted(overlap)}")
    # A Tier 0 class must never carry an outbound verb.
    outbound = {"compose_message", "export_via_share_sheet", "open_url", "search"}
    for name, c in classes.items():
        if c["tier"] == 0:
            bad = {a["verb"] for slot in ("primary", "secondary") for a in c[slot]} & outbound
            if bad:
                errors.append(f"{name}: Tier 0 인데 외부 전송 동사 {sorted(bad)}")
    return errors


def main() -> int:
    ap = argparse.ArgumentParser(description="엑셀 -> actions.json")
    ap.add_argument("--check", action="store_true",
                    help="엑셀이 JSON보다 새로우면 실패 (빌드 전 확인용)")
    ap.add_argument("--dry-run", action="store_true", help="리포트만, 쓰지 않음")
    args = ap.parse_args()

    if not XLSX.is_file():
        sys.exit(f"엑셀이 없습니다: {XLSX}")

    if args.check:
        if not OUT_JSON.is_file():
            sys.exit(f"actions.json 이 없습니다. 먼저 생성하세요:\n  python {Path(__file__).name}")
        digest = hashlib.sha256(XLSX.read_bytes()).hexdigest()
        recorded = json.loads(OUT_JSON.read_text(encoding="utf-8")).get("sourceSha256")
        if digest != recorded:
            sys.exit(
                "엑셀이 actions.json 보다 최신입니다 (내용 해시 불일치).\n"
                f"  python {Path(__file__).name}   으로 다시 생성하세요."
            )
        print("actions.json 이 엑셀과 일치합니다.")
        return 0

    warnings: list[str] = []
    wb = openpyxl.load_workbook(XLSX, data_only=True)

    verbs = load_verbs(wb)
    classes, unmapped = load_classes(wb, set(verbs))
    blocking = load_blocking(wb)
    signals = load_signals(wb)
    apply_review_adjustments(classes, verbs, warnings.append)

    errors = validate(classes, verbs, blocking)

    n_primary_ok = sum(1 for c in classes.values() if c["primary"])
    n_no_primary = sum(1 for c in classes.values() if not c["primary"])
    n_no_actions = sum(1 for c in classes.values() if not c["primary"] and not c["secondary"])
    n_unmapped_primary = sum(1 for c in classes.values() for u in c["unmapped"] if u["slot"] == "primary")
    n_unmapped_secondary = sum(1 for c in classes.values() for u in c["unmapped"] if u["slot"] == "secondary")

    print(f"클래스 {len(classes)}  차단 {len(blocking)}  동사 {len(verbs)}  신호 {len(signals)}")
    print(f"주 액션 매핑 {n_primary_ok}/{len(classes)}   미매핑: 주 {n_unmapped_primary} · 부 {n_unmapped_secondary}")
    print(f"주 액션 없는 클래스 {n_no_primary}개 · 액션이 하나도 없는 클래스 {n_no_actions}개")
    print("  (유니버설 액션은 항상 깔리므로 빈 화면이 되지는 않습니다)")
    for w in warnings:
        print(f"  경고: {w}")
    for e in errors:
        print(f"  오류: {e}")
    if errors:
        return 1

    catalog = {
        "schemaVersion": SCHEMA_VERSION,
        "generatedFrom": str(XLSX.relative_to(REPO)),
        "generatedAt": datetime.now(timezone.utc).isoformat(timespec="seconds"),
        "sourceSha256": hashlib.sha256(XLSX.read_bytes()).hexdigest(),
        "baseScores": {
            "primary": SCORE_PRIMARY,
            "secondaryFirst": SCORE_SECONDARY_FIRST,
            "secondaryRest": SCORE_SECONDARY_REST,
            "universal": SCORE_UNIVERSAL,
        },
        "universalActions": [
            {"verb": v, "baseScore": SCORE_UNIVERSAL} for v in UNIVERSAL_VERBS
        ],
        "verbs": verbs,
        "classes": classes,
        "blockingClasses": blocking,
        "rankingSignals": signals,
    }

    if args.dry_run:
        print("(dry-run: 쓰지 않음)")
        return 0

    RESOURCES.mkdir(parents=True, exist_ok=True)
    OUT_JSON.write_text(json.dumps(catalog, ensure_ascii=False, indent=2) + "\n",
                        encoding="utf-8")
    unmapped_rows = sum(len(c["unmapped"]) for c in classes.values())
    with UNMAPPED_CSV.open("w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=["class", "slot", "display", "verb"])
        w.writeheader()
        # Written from the adjusted classes, not the raw parse, so a slot the
        # review already settled does not reappear as work to do.
        for name, c in classes.items():
            for u in c["unmapped"]:
                w.writerow({"class": name, "slot": u["slot"], "display": u["display"], "verb": ""})

    print(f"\n생성: {OUT_JSON.relative_to(REPO)}  ({OUT_JSON.stat().st_size // 1024} KB)")
    print(f"미매핑 리포트: {UNMAPPED_CSV.relative_to(REPO)}  ({unmapped_rows}행)")
    print("  verb 열을 채우고 엑셀에 반영한 뒤 다시 생성하세요. 추측해서 채우지 않았습니다.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
